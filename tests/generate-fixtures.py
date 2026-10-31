"""Create synthetic local fixtures; never uses real warehouse data."""
from pathlib import Path
from openpyxl import Workbook
out=Path(__file__).resolve().parent/'fixtures';out.mkdir(exist_ok=True)
for name,count in [('small.xlsx',3),('large.xlsx',7500)]:
 book=Workbook();sheet=book.active;sheet.title='کالاها'
 sheet.append(['صورت شمارش آزمایشی']);sheet.append(['موجودی سیستم','نام کالا','شماره تگ','کد کالا'])
 for i in range(count):
  sheet.append([12.5+i,f'کالای آزمایشی {i+1}',str(i+5),i+17])
  sheet.cell(i+3,4).number_format='00000'
 book.create_sheet('توضیحات').append(['این فایل فقط برای آزمون است.'])
 book.save(out/name)
print('Created small.xlsx and large.xlsx (7500 items)')
